"""Synthetic regression checks for continuous-window diagnostics."""
import csv
import tempfile
import unittest
from pathlib import Path

from audit_ved_windows import audit, segment_points
from openpyxl import Workbook


class WindowAuditTest(unittest.TestCase):
    def test_invalid_rows_and_timestamp_gaps_split_segments(self):
        points = [
            (0, {"raw_rate": 1}),
            (2_000, {"raw_rate": 0}),
            (4_000, {"raw_rate": None}),
            (6_000, {"raw_rate": 0}),
            (20_000, {"raw_rate": 0}),
        ]
        self.assertEqual(segment_points(points, "raw_rate", 5_000),
                         [(2_000, 2, False), (0, 1, True), (0, 1, True)])
        self.assertEqual(segment_points(points, "raw_rate", 30_000),
                         [(2_000, 2, False), (14_000, 2, True)])

    def test_duplicates_break_continuity(self):
        points = [(0, {"maf": 1}), (1_000, {"maf": 2}), (1_000, {"maf": 3})]
        self.assertEqual(segment_points(points, "maf", 5_000),
                         [(1_000, 2, False), (0, 1, False)])

    def test_aggregate_has_no_identifiers_and_counts_thresholds(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            static = root / "static.xlsx"
            book = Workbook()
            book.active.append(["VehId", "Vehicle Type"])
            book.active.append([987654321, "ICE"])
            book.save(static)
            dynamic = root / "dynamic.csv"
            headers = ["VehId", "Trip", "Timestamp(ms)", "Fuel Rate[L/hr]", "MAF[g/sec]",
                       "Short Term Fuel Trim Bank 1[%]", "Long Term Fuel Trim Bank 1[%]"]
            with dynamic.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                for stamp in range(0, 301_000, 5_000):
                    writer.writerow([987654321, "private-trip", stamp, "", 2, 0, 0])
            result = audit(dynamic, [static])
            encoded = str(result)
            self.assertNotIn("987654321", encoded)
            self.assertNotIn("private-trip", encoded)
            counts = result["by_powertrain"]["ICE"]["channels"]["maf_bank1_trims"]["max_gap_5s"]
            self.assertEqual(counts["segments_at_least_300s"], 1)
            self.assertEqual(counts["vehicles_with_segments_at_least_300s"], 1)

    def test_missing_timestamp_excludes_the_group_from_continuity(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            static = root / "static.xlsx"
            book = Workbook()
            book.active.append(["VehId", "Vehicle Type"])
            book.active.append([1, "PHEV"])
            book.save(static)
            dynamic = root / "dynamic.csv"
            headers = ["VehId", "Trip", "Timestamp(ms)", "Fuel Rate[L/hr]", "MAF[g/sec]",
                       "Short Term Fuel Trim Bank 1[%]", "Long Term Fuel Trim Bank 1[%]"]
            with dynamic.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                writer.writerow([1, "a", 0, 1, 2, 0, 0])
                writer.writerow([1, "a", "", 1, 2, 0, 0])
                writer.writerow([1, "a", 600_000, 1, 2, 0, 0])
            result = audit(dynamic, [static])
            self.assertEqual(result["by_powertrain"]["PHEV"]["quality"]
                             ["groups_with_missing_timestamp"], 1)
            self.assertEqual(result["by_powertrain"]["PHEV"]["channels"]
                             ["raw_rate"]["max_gap_30s"]["segments"], 0)

    def test_source_order_reversal_is_reported_and_sources_are_hashed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            static = root / "static.xlsx"
            book = Workbook()
            book.active.append(["VehId", "Vehicle Type"])
            book.active.append([1, "ICE"])
            book.save(static)
            dynamic = root / "dynamic.csv"
            headers = ["VehId", "Trip", "Timestamp(ms)", "Fuel Rate[L/hr]", "MAF[g/sec]",
                       "Short Term Fuel Trim Bank 1[%]", "Long Term Fuel Trim Bank 1[%]"]
            with dynamic.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                writer.writerow([1, "a", 5_000, 1, 2, 0, 0])
                writer.writerow([1, "a", 0, 1, 2, 0, 0])
            result = audit(dynamic, [static])
            self.assertEqual(result["by_powertrain"]["ICE"]["quality"]
                             ["groups_with_source_order_reversal"], 1)
            self.assertEqual(len(result["dynamic_sha256"]), 64)
            self.assertEqual(len(result["static_sources"][0]["sha256"]), 64)


if __name__ == "__main__":
    unittest.main()
